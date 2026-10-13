#!/usr/bin/env python3
"""Verify every slide citation in the Day 2 practice materials against the Day 2 deck.

    python build/check_slide_refs.py "<path>/Genomic Data Processing · 10 13 Molecular tools for genomics.pptx"

The deck is not part of this repository, so pass its path. Requires python-pptx.

Scans README.md and both notebooks (markdown and code) for three citation forms:
  * prose:            "slide 24", "slides 14-16", "practice slide 19"
  * number + title:   6 "Measure the molecule before you sequence"   (README mapping table)
                      "Cohort QC: is every sample who it claims to be?" (slide 24)   (notebooks)
                      19 · Mini practice 1 · Is this sample fit for this platform?
Every cited number must exist in the deck. A number cited together with a title must carry that
slide's title. A number cited without a title must have an expected title recorded in EXPECT_SLIDE;
a cited slide with no expectation is reported as UNVERIFIED and fails the run, so silence never
means "probably fine". Exits non-zero on any problem.
"""
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent.parent
SOURCES = ["README.md", "practice1_sample_to_platform.ipynb", "practice2_cohort_qc.ipynb"]

# Expected title for every slide cited by number. Keep in step with the text: if you cite a new
# slide, record what it should say.
EXPECT_SLIDE = {
    2: "What is a genomics sample?",
    6: "Measure the molecule before you sequence",
    7: "Metadata: the sample's passport",
    14: "Illumina: short, accurate, massively parallel",
    15: "PacBio HiFi: long and accurate",
    16: "Oxford Nanopore: the longest reads",
    17: "Side by side: which platform for which question?",
    18: "Match the sample to the platform",
    19: "Mini practice 1 · Is this sample fit for this platform?",
    24: "Cohort QC: is every sample who it claims to be?",
    25: "What a cohort can and cannot tell you",
    26: "Mini practice 2 · Cohort QC: who does not belong?",
    29: "Take-home: what is a genomics sample?",
}

PROSE = re.compile(r"\bslides?\s+(\d{1,2}(?:\s*[-–]\s*\d{1,2})?(?:\s*,\s*\d{1,2})*)", re.I)
NUM_TITLE = re.compile(r"(?<![\d.])(\d{1,2}) \"([^\"]+)\"")                 # 6 "Title"
TITLE_NUM = re.compile(r"\"([^\"]+)\" \(slide (\d{1,2})\)")                 # "Title" (slide 6)
PRACTICE = re.compile(r"(?<![\d.])(\d{1,2}) · (Mini practice [12] · [^|\n*]+?)\s*(?:\||$)", re.M)


def norm(t):
    return re.sub(r"\s+", " ", t.replace("’", "'")).strip()


def expand(spec):
    out = []
    for part in spec.split(","):
        part = part.strip()
        m = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", part)
        if m:
            out += list(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def deck_titles(path):
    prs = Presentation(path)
    titles = []
    for s in prs.slides:
        ts = [norm(sh.text_frame.text) for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(ts[1] if len(ts) > 1 else (ts[0] if ts else ""))
    return titles


def source_text(name):
    p = HERE / name
    if p.suffix == ".ipynb":
        nb = json.loads(p.read_text())
        return "\n".join("".join(c["source"]) for c in nb["cells"])
    return p.read_text()


def main(deck_path):
    titles = deck_titles(deck_path)
    total = len(titles)
    problems, n_pairs, n_prose = [], 0, 0
    for name in SOURCES:
        txt = norm_lines = source_text(name)
        flat = norm(txt)
        pairs = [(int(n), t) for n, t in NUM_TITLE.findall(flat)]
        pairs += [(int(n), t) for t, n in TITLE_NUM.findall(flat)]
        pairs += [(int(n), t) for n, t in PRACTICE.findall(norm_lines)]
        for n, t in pairs:
            if not 1 <= n <= total:
                problems.append(f"{name}: slide {n} is outside the deck (1-{total}) [{t}]")
            elif norm(t) not in titles[n - 1]:
                problems.append(f"{name}: cites slide {n} as '{t}', but slide {n} is '{titles[n - 1]}'")
            else:
                n_pairs += 1
        for m in PROSE.finditer(flat):
            for n in expand(m.group(1)):
                if not 1 <= n <= total:
                    problems.append(f"{name}: slide {n} is outside the deck (1-{total}) [{m.group(0)}]")
                    continue
                want = EXPECT_SLIDE.get(n)
                if want is None:
                    problems.append(f"UNVERIFIED {name}: slide {n} is cited [{m.group(0)}] but EXPECT_SLIDE has "
                                    f"no title for it; the deck says '{titles[n - 1]}'")
                elif want != titles[n - 1]:
                    problems.append(f"{name}: slide {n} should be '{want}' but the deck says '{titles[n - 1]}'")
                else:
                    n_prose += 1
    for n, want in EXPECT_SLIDE.items():
        if n > total or titles[n - 1] != want:
            problems.append(f"EXPECT_SLIDE: slide {n} should be '{want}' but the deck says "
                            f"'{titles[n - 1] if n <= total else 'missing'}'")
    print(f"deck: {total} slides | number+title citations checked: {n_pairs} | prose citations checked: {n_prose}")
    for p in problems:
        print("PROBLEM", p)
    if problems:
        sys.exit(1)
    print("all slide citations match the deck")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
