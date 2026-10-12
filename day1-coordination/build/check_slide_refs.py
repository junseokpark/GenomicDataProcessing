#!/usr/bin/env python3
"""Verify every slide reference in the practice materials against a deck.

    python build/check_slide_refs.py ../Lecture-Day1-Jun.pptx

Scans three places, because they use different notations and it is easy to miss one:
  * the Slides column of the mapping table in README.md   (bare ranges: "59-62, 84")
  * prose in README.md                                    ("slide 82", "slides 51-54")
  * the notebook's markdown and code                      (same prose forms)

Exits non-zero if any reference is out of range or lands outside its claimed section.
"""
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent.parent

# README mapping-table section -> eyebrow keywords that are acceptable for its slides
EXPECT = {
    "1 Reference / history":  ("GUIDED Q&A", "PRACTICE"),
    "2 hg19 coordinates":     ("COORDINATES", "GUIDED Q&A", "PRACTICE"),
    "3 VCF semantics":        ("VCF SEMANTICS", "GUIDED Q&A", "PRACTICE"),
    "4 hg19 → hg38 mapping":  ("GUIDED Q&A", "PRACTICE"),
    "5 Assembly-matched IGV": ("EVIDENCE", "GUIDED Q&A", "PRACTICE"),
    "6 Read realignment":     ("ALIGNMENT", "GUIDED Q&A", "PRACTICE"),
    "7 Batch and QC":         ("PROCESSING PRINCIPLES", "SCALE", "PRACTICE"),
    "8 Evidence card":        ("YOUR PRESENTATION", "PRACTICE", "PRESENTATIONS"),
    "Appendix":               ("ACTIVITY",),
}
PROSE = re.compile(
    r"slides?\s*((?:\d{1,3}(?:\s*[-–—]\s*\d{1,3})?)(?:\s*,\s*\d{1,3}(?:\s*[-–—]\s*\d{1,3})?)*)",
    re.I)


def expand(spec):
    out = []
    for part in spec.split(","):
        part = part.strip()
        m = re.fullmatch(r"(\d+)\s*[-–—]\s*(\d+)", part)
        if m:
            out += list(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def deck_text(prs, n):
    ts = [sh.text_frame.text.strip().replace("\n", " ") for sh in prs.slides[n - 1].shapes
          if sh.has_text_frame and sh.text_frame.text.strip()]
    return (ts[0] if ts else ""), (ts[1] if len(ts) > 1 else "")


def main(deck_path):
    prs = Presentation(deck_path)
    total = len(prs.slides)
    readme = (HERE / "README.md").read_text()
    nb = json.dumps(json.loads((HERE / "genomic_data_processing_practice.ipynb").read_text()))

    refs = []                                            # (source, slide, context)
    for cells in re.findall(r"^\|([^|]+)\|([^|]+)\|([^|]+)\|\s*$", readme, re.M):
        section, spec = cells[0].strip(), cells[1].strip()
        for n in expand(spec):
            refs.append(("README table", n, section))
    for src, txt in (("README prose", readme), ("notebook", nb)):
        for m in PROSE.finditer(txt):
            for n in expand(m.group(1)):
                refs.append((src, n, m.group(0)[:40]))

    problems = []
    for src, n, ctx in refs:
        if not 1 <= n <= total:
            problems.append(f"{src}: slide {n} is outside the deck (1-{total}) [{ctx}]")
            continue
        if src == "README table":
            expected = EXPECT.get(ctx)
            eyebrow, title = deck_text(prs, n)
            if expected and not any(k in eyebrow for k in expected):
                problems.append(
                    f"{src}: row '{ctx}' cites slide {n}, but that slide is "
                    f"'{eyebrow} | {title}' (expected one of {expected})")

    by_src = {}
    for src, n, _ in refs:
        by_src.setdefault(src, set()).add(n)
    print(f"deck: {Path(deck_path).name}  ({total} slides)")
    for src in sorted(by_src):
        print(f"  {src:14s} {len(by_src[src]):3d} distinct slides")
    print(f"  {'TOTAL':14s} {len(refs):3d} reference instances, "
          f"{len({n for _, n, _ in refs})} distinct slides")

    missing_footer = [i for i in range(1, total + 1)
                      if not any(re.fullmatch(r"\d+\s*/\s*\d+", sh.text_frame.text.strip())
                                 for sh in prs.slides[i - 1].shapes if sh.has_text_frame)]
    if missing_footer:
        print(f"  note: slides with no footer page number: {missing_footer}")

    if problems:
        print("\nPROBLEMS:")
        for p in problems:
            print("  -", p)
        return 1
    print("\nall slide references resolve correctly")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
